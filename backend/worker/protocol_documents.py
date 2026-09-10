"""Material parsing and private content-addressed storage for protocol extraction."""
import hashlib
import io
import json
import os
import re
import zipfile
from pathlib import Path
from uuid import uuid4

import pdfplumber

MAX_FILE_BYTES = 30 * 1024 * 1024
MAX_TEXT_CHARS = 20_000_000
MAX_PAGES = 300
TEXT_FORMATS = {'.pdf', '.txt', '.md', '.docx', '.xlsx', '.csv', '.tsv', '.zip'}


def document_root():
    return Path(os.environ.get('PROTOCOL_DIR', 'data/protocols')).resolve()


def parse_document(content: bytes, name: str):
    try:
        return _parse_document(content, name)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('材料损坏、加密或无法解析，请检查原文件') from exc


def _parse_document(content: bytes, name: str):
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError('材料为空或超过 30 MB')
    suffix = Path(name).suffix.lower()
    warnings = []
    if suffix == '.pdf':
        if not content.startswith(b'%PDF-'):
            raise ValueError('文件内容不是 PDF')
        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                if len(pdf.pages) > MAX_PAGES:
                    raise ValueError('PDF 超过 300 页，请拆分材料')
                pages = []
                size = 0
                for number, page in enumerate(pdf.pages, 1):
                    text = page.extract_text(use_text_flow=True) or ''
                    size += len(text)
                    if size > MAX_TEXT_CHARS:
                        raise ValueError('提取文本超过容量，请拆分材料')
                    if len(text.strip()) < 30:
                        warnings.append(f'第 {number} 页文本不足，可能需要 OCR 或人工检查')
                    pages.append({'page': number, 'text': text})
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('PDF 损坏、加密或无法解析') from exc
    elif suffix == '.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        _check_archive(content)
        doc = Document(io.BytesIO(content))
        blocks = ['Word document text and tables (not physical page numbers).']
        for child in doc.element.body:
            tag = child.tag.rsplit('}', 1)[-1]
            if tag == 'p':
                blocks.append(Paragraph(child, doc).text)
            elif tag == 'tbl':
                blocks.extend('\t'.join(cell.text.replace('\n', ' | ') for cell in row.cells)
                              for row in Table(child, doc).rows)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            from xml.etree import ElementTree as ET
            for part in ('word/footnotes.xml', 'word/endnotes.xml'):
                if part in archive.namelist():
                    root = ET.fromstring(archive.read(part))
                    blocks.append(part + '\n' + '\n'.join(node.text or '' for node in root.iter()
                                                           if node.tag.rsplit('}', 1)[-1] == 't'))
            if any(name.startswith('word/media/') for name in archive.namelist()):
                warnings.append('Word 含内嵌图像；其中的方法或配方需人工检查或 OCR')
        text = '\n\n'.join(blocks)
        pages = [{'page': 1, 'text': text}]
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        _check_archive(content)
        readable = content
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if b'http://purl.oclc.org/ooxml/spreadsheetml/main' in archive.read('xl/workbook.xml'):
                converted = io.BytesIO()
                with zipfile.ZipFile(converted, 'w', zipfile.ZIP_DEFLATED) as target:
                    for member in archive.infolist():
                        data = archive.read(member)
                        if member.filename.endswith(('.xml', '.rels')):
                            data = data.replace(b'http://purl.oclc.org/ooxml/spreadsheetml/main', b'http://schemas.openxmlformats.org/spreadsheetml/2006/main')
                            data = data.replace(b'http://purl.oclc.org/ooxml/officeDocument/relationships', b'http://schemas.openxmlformats.org/officeDocument/2006/relationships')
                        target.writestr(member.filename, data)
                readable = converted.getvalue()
                warnings.append('Strict OOXML 表格按兼容命名空间读取，原始文件保持不变')
        workbook = load_workbook(io.BytesIO(readable), read_only=True, data_only=False)
        pages, size = [], 0
        try:
            for index, sheet in enumerate(workbook, 1):
                lines = [f'Worksheet: {sheet.title} (section {index}, not a PDF page)']
                for row in sheet.iter_rows(values_only=True):
                    if all(cell is None for cell in row):
                        continue
                    line = '\t'.join('' if cell is None else str(cell).replace('\n', ' | ') for cell in row)
                    size += len(line)
                    if size > MAX_TEXT_CHARS:
                        raise ValueError('表格文本超过容量，请拆分包含方法或配方的工作表')
                    lines.append(line)
                pages.append({'page': index, 'text': '\n'.join(lines)})
        finally:
            workbook.close()
    elif suffix == '.zip':
        _check_archive(content)
        pages = []
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            for member in archive.infolist():
                if member.is_dir():
                    continue
                if Path(member.filename).suffix.lower() not in TEXT_FORMATS - {'.zip'}:
                    warnings.append(f'{member.filename}: 非可搜索文本附件，需人工检查或 OCR')
                    continue
                try:
                    nested, issues = parse_document(archive.read(member), member.filename)
                    for page in nested:
                        pages.append({'page': len(pages) + 1,
                                      'text': f'Archive member: {member.filename}; source page/section: {page["page"]}\n\n{page["text"]}'})
                    warnings.extend(f'{member.filename}: {issue}' for issue in issues)
                except ValueError as exc:
                    warnings.append(f'{member.filename}: {exc}')
    elif suffix in {'.txt', '.md', '.csv', '.tsv'}:
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('文本文件须使用 UTF-8 编码') from exc
        if len(text) > MAX_TEXT_CHARS:
            raise ValueError('文本超过容量，请拆分材料')
        pages = [{'page': 1, 'text': text}]
    else:
        raise ValueError('当前支持 PDF、TXT、MD、DOCX、XLSX、CSV、TSV 和 ZIP 补充材料')
    if sum(len(p['text']) for p in pages) > MAX_TEXT_CHARS:
        raise ValueError('提取文本超过容量，请拆分材料')
    if sum(len(p['text'].strip()) for p in pages) < 50:
        raise ValueError('材料没有足够的可搜索文本，请上传 OCR 后的 PDF 或对应文本')
    return pages, warnings


def _check_archive(content):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(member.file_size for member in archive.infolist()) > MAX_FILE_BYTES * 5:
                raise ValueError('压缩文件展开后过大，请拆分方法或配方附件')
    except zipfile.BadZipFile as exc:
        raise ValueError('压缩文件或 Office 文档损坏') from exc


def save_document(content, name, item_id):
    pages, warnings = parse_document(content, name)
    sha256 = hashlib.sha256(content).hexdigest()
    directory = document_root() / str(item_id)
    directory.mkdir(parents=True, exist_ok=True)
    destination = directory / f'{sha256}{Path(name).suffix.lower()}'
    temporary = directory / f'{uuid4().hex}.partial'
    try:
        temporary.write_bytes(content)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)
    return {'name': Path(name).name[:256], 'path': str(destination), 'sha256': sha256,
            'pages_json': json.dumps(pages, ensure_ascii=False),
            'warnings_json': json.dumps(warnings, ensure_ascii=False)}


def document_context(documents):
    return {d.id: {'name': d.name, 'reference': d.reference, 'kind': d.kind,
                   'pages': json.loads(d.pages_json)} for d in documents}


def material_chunks(documents, max_chars=18000):
    """Chunk every included source section; exclusions remain in the coverage manifest."""
    chunks, current, length = [], [], 0
    sections, _ = protocol_pages(documents)
    for section in sections:
        text = section['text']
        # Preserve overlap without dropping any included source text.
        for offset in range(0, max(1, len(text)), max_chars - 1000):
            part = text[offset:offset + max_chars]
            if current and length + len(part) > max_chars:
                chunks.append(current); current, length = [], 0
            current.append(dict(section, text=part)); length += len(part)
    if current:
        chunks.append(current)
    return chunks


def material_manifest_path(item_id):
    return document_root() / str(item_id) / 'materials-manifest.json'


def read_material_manifest(item_id):
    path = material_manifest_path(item_id)
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else {}


def protocol_pages(documents):
    """Methods first; all supplementary narrative/recipes, with indexed data tables."""
    has_methods = any(doc.kind == 'methods' for doc in documents)
    sections, coverage = [], []
    for doc in documents:
        pages = json.loads(doc.pages_json)
        if has_methods and doc.kind == 'main':
            coverage.append({'document_id': doc.id, 'name': doc.name, 'kind': doc.kind, 'reference': doc.reference, 'included': False,
                             'reason': '原始全文保留供核对；同篇 Methods 已单独纳入'})
            continue
        for page in pages:
            text = page['text']
            # Omit only explicit gene/statistical result tables without recipe clues.
            # Scan the whole section for dosing/methods before classifying it as data.
            table = 'Worksheet:' in text or Path(doc.name).suffix.lower() in {'.csv', '.tsv'}
            data_header = re.search(r'Ensembl\s*ID|log2\s*Fold\s*Change|avg_log2?FC|adj[._ ]?pval|Fold Enrichment|\bpadj\b|\blogFC\b|enrichmentScore|Motif Name\s+Consensus|gene_short_name|module\s+supermodule|Combined Score\s+Genes', text[:6000], re.I)
            recipe = re.search(r'\bprotocol\b|\bmethods\b|\bRPMI\b|\bDMEM\b|\bmTeSR\b|\bconcentration\b|\bdose\b|\bCHIR(?:99021)?\b|\bActivin\s+A\b|\d\s*(?:[nuμµm]M|ng/m[lL]|[μµu]g/m[lL])\b', text, re.I)
            include = not (table and data_header and not recipe)
            coverage.append({'document_id': doc.id, 'name': doc.name, 'kind': doc.kind, 'reference': doc.reference, 'page': page['page'],
                             'included': include, 'characters': len(text),
                             'reason': '纳入 Methods / 补充材料' if include else '已扫描：基因或统计结果表，未发现方法或配方字段'})
            if include:
                sections.append({'document_id': doc.id, 'filename': doc.name, 'kind': doc.kind,
                                 'reference': doc.reference, 'page': page['page'], 'text': text})
    return sections, coverage
