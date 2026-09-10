import io
import json
import pytest

from backend.worker.protocol_documents import parse_document


def test_methods_include_nested_recipe_table_and_only_cited_references():
    from backend.worker.protocol_sources import parse_article_xml
    xml = b'''<article xmlns:xlink="http://www.w3.org/1999/xlink"><body>
      <sec><title>Results</title><p>Result unrelated to the recipe.</p></sec>
      <sec sec-type="methods"><title>Materials and Methods</title>
        <sec><title>Cell differentiation</title><p>As described previously <xref ref-type="bibr" rid="R17">17</xref>, cells received CHIR99021.</p>
          <table-wrap><caption><p>Stage recipe</p></caption><table><tr><th>Agent</th><th>Concentration</th></tr><tr><td>CHIR99021</td><td>3 uM</td></tr></table></table-wrap>
        </sec></sec></body><back><ref-list>
      <ref id="R17"><label>17</label><element-citation><article-title>Original recipe</article-title><pub-id pub-id-type="doi">10.1/recipe</pub-id></element-citation></ref>
      <ref id="R18"><label>18</label><mixed-citation>Unrelated reference</mixed-citation></ref>
      </ref-list><supplementary-material><caption><p>Supplementary methods</p></caption><media xlink:href="methods.docx"/></supplementary-material></back></article>'''
    result = parse_article_xml(xml)
    assert 'CHIR99021\t3 uM' in result['methods']
    assert 'Cell differentiation' in result['methods']
    assert 'Result unrelated' not in result['methods']
    assert [x['id'] for x in result['references']] == ['R17']
    assert result['supplements'][0]['href'] == 'methods.docx'
    assert result['supplements'][0]['caption'] == 'Supplementary methods'


def test_docx_supplement_keeps_paragraphs_and_recipe_table():
    from docx import Document
    doc = Document();doc.add_paragraph('Supplementary Methods: differentiation starts on day 0.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0,0).text='Factor';table.cell(0,1).text='Dose'
    table.cell(1,0).text='BMP4';table.cell(1,1).text='20 ng/mL'
    data=io.BytesIO();doc.save(data)
    pages,warnings=parse_document(data.getvalue(),'supplement.docx')
    text='\n'.join(p['text'] for p in pages)
    assert 'day 0' in text and 'BMP4\t20 ng/mL' in text


def test_xlsx_supplement_retains_all_worksheets_and_cell_values():
    from openpyxl import Workbook
    wb=Workbook();wb.active.title='Stage 1';wb.active.append(['CHIR99021',3,'uM'])
    wb.create_sheet('Stage 2').append(['IWR1',5,'uM'])
    data=io.BytesIO();wb.save(data)
    pages,warnings=parse_document(data.getvalue(),'recipes.xlsx')
    assert len(pages)==2
    assert 'Stage 1' in pages[0]['text'] and 'CHIR99021\t3\tuM' in pages[0]['text']
    assert 'Stage 2' in pages[1]['text'] and 'IWR1\t5\tuM' in pages[1]['text']


def test_zip_supplement_preserves_member_and_page_identity():
    import zipfile
    data=io.BytesIO()
    with zipfile.ZipFile(data,'w') as z:
        z.writestr('supplement/methods.txt','Supplementary protocol: human iPSC differentiation with BMP4 at 20 ng/mL, from day 0 to day 2.')
        z.writestr('supplement/movie.mp4',b'not text')
    pages,warnings=parse_document(data.getvalue(),'supplement.zip')
    assert 'supplement/methods.txt' in pages[0]['text']
    assert 'BMP4' in pages[0]['text']
    assert any('movie.mp4' in warning for warning in warnings)


def test_strict_ooxml_supplement_preserves_values():
    import zipfile
    from openpyxl import Workbook
    wb=Workbook();wb.active.title='Recipe';wb.active.append(['Differentiation protocol CHIR99021',3,'uM from D0 to D2'])
    data=io.BytesIO();wb.save(data);strict=io.BytesIO()
    with zipfile.ZipFile(data) as src, zipfile.ZipFile(strict,'w') as out:
        for name in src.namelist():
            content=src.read(name)
            if name.endswith(('.xml','.rels')):
                content=content.replace(b'http://schemas.openxmlformats.org/spreadsheetml/2006/main',b'http://purl.oclc.org/ooxml/spreadsheetml/main').replace(b'http://schemas.openxmlformats.org/officeDocument/2006/relationships',b'http://purl.oclc.org/ooxml/officeDocument/relationships')
            out.writestr(name,content)
    pages,warnings=parse_document(strict.getvalue(),'strict.xlsx')
    assert len(pages)==1 and 'CHIR99021\t3\tuM' in pages[0]['text']
    assert any('Strict' in w for w in warnings)


def test_coverage_keeps_late_recipe_in_result_table_and_archives_main_pdf():
    from types import SimpleNamespace
    from backend.worker.protocol_documents import protocol_pages
    def doc(id,kind,name,text):
        return SimpleNamespace(id=id,kind=kind,name=name,reference='source',pages_json=json.dumps([{'page':1,'text':text}]))
    methods=doc(2,'methods','Methods.md','Methods CHIR99021 at 3 uM')
    docs=[doc(1,'main','main.pdf','full paper'),methods,
          doc(3,'supplement','data.xlsx','Worksheet: Genes\nEnsembl ID\tpadj\n'+ 'ENSG1\t0.1\n'*1000),
          doc(4,'supplement','conditions.xlsx','Worksheet: Genes\nEnsembl ID\tpadj\n'+'ENSG1\t0.1\n'*1000+'CHIR99021 concentration 3 uM')]
    pages,coverage=protocol_pages(docs)
    assert {p['document_id'] for p in pages}=={2,4}
    assert len(coverage)==4 and all(c.get('reason') for c in coverage)


def test_jove_protocol_section_is_a_method_source():
    from backend.worker.protocol_sources import parse_article_xml
    xml=b'<article><body><sec><title>Protocol</title><sec><title>PGCLC induction</title><p>Add BMP4 at 500 ng/mL from day zero.</p></sec></sec></body></article>'
    result=parse_article_xml(xml)
    assert result['method_titles']==['Protocol']
    assert 'BMP4 at 500 ng/mL' in result['methods']
